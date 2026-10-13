"""Dinamik Takvim Excel dosyasından index.html üretir.

Kullanım:  python3 tools/excel_to_html.py Dinamik_Takvim.xlsx [cikti.html]
"""
import openpyxl,sys,json,datetime as dt,pathlib
wb=openpyxl.load_workbook(sys.argv[1])
def d(v): return v.date().isoformat() if hasattr(v,'year') else None
def t(v):
    if v is None: return ""
    if isinstance(v,dt.time): return v.strftime("%H:%M")
    s=str(v).strip().replace(".",":")
    return s
def s(v): return "" if v is None else str(v).strip()
ws=wb.worksheets[1]
tasks=[]
for r in range(6,ws.max_row+1):
    g=lambda c: ws.cell(r,c).value
    if not (g(13) or g(7)): continue
    tasks.append(dict(id=len(tasks)+1,b=s(g(2)),c=s(g(3)),d=s(g(4)),e=s(g(5)),f=s(g(6)),
        bas=d(g(7)),bit=d(g(8)),saatBas=t(g(9)),saatBit=t(g(10)),yaziTar=d(g(11)),yaziNo=s(g(12)),is_=s(g(13))))
ws=wb.worksheets[0]
daily=[];done={}
for r in range(5,ws.max_row+1):
    L=s(ws.cell(r,12).value)
    if L and L not in daily: daily.append(L)
    f=d(ws.cell(r,6).value)
    if f and ws.cell(r,5).value=='T': done.setdefault(f,[]).append(daily.index(L))
for k in done: done[k]=sorted(set(done[k]))
out=dict(baslangic=d(wb.worksheets[1]['E1'].value),gunlukBaslangic=d(wb.worksheets[0]['E1'].value),isler=tasks,gunlukIsler=daily,gunlukYapilan=done)
for x in tasks: x['is']=x.pop('is_')
kok=pathlib.Path(__file__).resolve().parent.parent
sablon=(kok/"src"/"takvim.template.html").read_text(encoding="utf-8")
veri=json.dumps(out,ensure_ascii=False,separators=(',',':')).replace("</","<\\/")
html=sablon.replace("/*__VERI__*/null",veri)
hedef=pathlib.Path(sys.argv[2]) if len(sys.argv)>2 else kok/"index.html"
hedef.write_text(html,encoding="utf-8")
print(f"{hedef}: {len(tasks)} iş, {len(daily)} günlük iş, {len(done)} gün işaretli")
